"""
文档处理模块 - 从各种文件格式中提取文本

Document processing module for extracting text from various file formats.

核心功能 (Core Features):
1. 多格式支持: PDF, DOCX, PPTX, EPUB, TXT, MD, arXiv
2. 文本提取: 从不同格式中提取纯文本内容
3. 元数据解析: 提取标题、作者、创建日期等信息
4. 文档管理: 保存、查询、删除文档及其元数据
"""
import json
import logging
from pathlib import Path
from typing import Optional, Dict, Any, List
from datetime import datetime

import fitz  # PyMuPDF - PDF处理库
import arxiv  # arXiv论文下载库
from docx import Document as DocxDocument  # Word文档处理
from pptx import Presentation  # PowerPoint处理
from ebooklib import epub  # EPUB电子书处理
from bs4 import BeautifulSoup  # HTML解析(用于EPUB)

from backend.models.schemas import Document, DocumentType
from backend.utils.file_utils import (
    get_file_hash,
    get_document_type,
    get_file_size,
    sanitize_filename
)
from config.settings import settings

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """
    文档处理器 - 处理和提取各种格式文档的文本

    Document Processor - Process and extract text from documents.

    主要职责 (Main Responsibilities):
    - 识别文档类型并调用相应的处理方法
    - 提取文本内容和元数据
    - 将处理结果保存为标准化格式
    - 管理文档的生命周期(创建、查询、删除)
    """

    def __init__(self):
        """
        初始化文档处理器

        Initialize document processor.

        创建必要的目录结构:
        - metadata_dir: 存储文档元数据(JSON)和提取的文本
        - document_dir: 存储原始文档文件
        """
        self.metadata_dir = Path(settings.metadata_path)
        self.document_dir = Path(settings.document_path)
        # 确保目录存在，不存在则创建
        self.metadata_dir.mkdir(parents=True, exist_ok=True)
        self.document_dir.mkdir(parents=True, exist_ok=True)

    def process_pdf(self, file_path: str) -> Dict[str, Any]:
        """
        从PDF文件提取文本和元数据

        Extract text and metadata from PDF file.

        处理流程 (Processing Flow):
        1. 使用PyMuPDF(fitz)打开PDF文件
        2. 逐页提取文本内容
        3. 提取PDF内置元数据(标题、作者等)
        4. 返回标准化的字典格式

        Args:
            file_path: PDF文件路径 (Path to PDF file)

        Returns:
            包含文本和元数据的字典 (Dictionary containing text and metadata):
            - text: 完整文本内容
            - num_pages: 页数
            - metadata: PDF元数据
            - title: 文档标题
        """
        try:
            # 打开PDF文档
            doc = fitz.open(file_path)
            text_content = []

            # 遍历所有页面，提取文本
            for page_num in range(len(doc)):
                page = doc[page_num]
                text_content.append(page.get_text())

            # 合并所有页面的文本，页面之间用双换行分隔
            full_text = "\n\n".join(text_content)

            # 提取PDF的内置元数据
            metadata = doc.metadata

            # 构建结果字典
            result = {
                "text": full_text,
                "num_pages": len(doc),
                "metadata": metadata,
                "title": metadata.get("title", Path(file_path).stem),  # 优先使用PDF标题，否则用文件名
            }

            # 关闭文档释放资源
            doc.close()
            return result

        except Exception as e:
            logger.error(f"Error processing PDF {file_path}: {e}")
            raise

    def process_text_file(self, file_path: str) -> Dict[str, Any]:
        """
        Extract text from plain text or markdown file.

        Args:
            file_path: Path to text file

        Returns:
            Dictionary containing text and metadata
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()

            return {
                "text": text,
                "num_pages": None,
                "metadata": {},
                "title": Path(file_path).stem,
            }
        except Exception as e:
            logger.error(f"Error processing text file {file_path}: {e}")
            raise

    def process_docx(self, file_path: str) -> Dict[str, Any]:
        """
        从Word文档提取文本和元数据

        Extract text and metadata from DOCX file.

        处理流程 (Processing Flow):
        1. 提取所有段落的文本
        2. 提取表格中的文本(用 | 分隔单元格)
        3. 提取文档属性(标题、作者、创建时间等)

        Args:
            file_path: DOCX文件路径 (Path to DOCX file)

        Returns:
            包含文本和元数据的字典 (Dictionary containing text and metadata)
        """
        try:
            doc = DocxDocument(file_path)
            text_content = []

            # 从段落中提取文本
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)

            # 从表格中提取文本
            # 每行用 | 分隔单元格，使结构更清晰
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells])
                    if row_text.strip():
                        text_content.append(row_text)

            # 合并所有文本
            full_text = "\n\n".join(text_content)

            # 提取文档核心属性
            core_properties = doc.core_properties
            metadata = {
                "title": core_properties.title or Path(file_path).stem,
                "author": core_properties.author,
                "subject": core_properties.subject,
                "created": str(core_properties.created) if core_properties.created else None,
                "modified": str(core_properties.modified) if core_properties.modified else None,
            }

            return {
                "text": full_text,
                "num_pages": None,  # DOCX没有固定页面概念
                "metadata": metadata,
                "title": metadata.get("title", Path(file_path).stem),
            }

        except Exception as e:
            logger.error(f"Error processing DOCX {file_path}: {e}")
            raise

    def process_pptx(self, file_path: str) -> Dict[str, Any]:
        """
        Extract text and metadata from PPTX file.

        Args:
            file_path: Path to PPTX file

        Returns:
            Dictionary containing text and metadata
        """
        try:
            prs = Presentation(file_path)
            text_content = []

            # Extract text from all slides
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = [f"=== Slide {slide_num} ==="]

                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())

                if len(slide_text) > 1:  # More than just the header
                    text_content.append("\n".join(slide_text))

            full_text = "\n\n".join(text_content)

            # Extract metadata
            core_properties = prs.core_properties
            metadata = {
                "title": core_properties.title or Path(file_path).stem,
                "author": core_properties.author,
                "subject": core_properties.subject,
                "created": str(core_properties.created) if core_properties.created else None,
                "modified": str(core_properties.modified) if core_properties.modified else None,
            }

            return {
                "text": full_text,
                "num_pages": len(prs.slides),
                "metadata": metadata,
                "title": metadata.get("title", Path(file_path).stem),
            }

        except Exception as e:
            logger.error(f"Error processing PPTX {file_path}: {e}")
            raise

    def process_epub(self, file_path: str) -> Dict[str, Any]:
        """
        Extract text and metadata from EPUB file with safe metadata access.

        Args:
            file_path: Path to EPUB file

        Returns:
            Dictionary containing text and metadata
        """
        def safe_get_epub_metadata(book, dc_type: str, default=None):
            """
            Safely extract metadata from EPUB file.

            Args:
                book: EPUB book object
                dc_type: Dublin Core metadata type
                default: Default value if metadata not found

            Returns:
                Metadata value or default
            """
            try:
                meta = book.get_metadata('DC', dc_type)
                if meta and len(meta) > 0 and len(meta[0]) > 0:
                    return meta[0][0]
            except (IndexError, TypeError, AttributeError) as e:
                logger.debug(f"Could not extract EPUB metadata '{dc_type}': {e}")
            return default

        try:
            book = epub.read_epub(file_path)
            text_content = []

            # Extract text from all items with error handling
            for item in book.get_items():
                try:
                    if item.get_type() == epub.ITEM_DOCUMENT:
                        # Parse HTML content
                        soup = BeautifulSoup(item.get_content(), 'html.parser')
                        text = soup.get_text(separator='\n', strip=True)
                        if text:
                            text_content.append(text)
                except Exception as e:
                    logger.warning(f"Failed to extract text from EPUB item: {e}")
                    continue

            full_text = "\n\n".join(text_content)

            # Extract metadata with safe access
            metadata = {
                "title": safe_get_epub_metadata(book, 'title', Path(file_path).stem),
                "author": safe_get_epub_metadata(book, 'creator'),
                "language": safe_get_epub_metadata(book, 'language'),
                "publisher": safe_get_epub_metadata(book, 'publisher'),
            }

            return {
                "text": full_text,
                "num_pages": None,  # EPUB doesn't have fixed pages
                "metadata": metadata,
                "title": metadata.get("title", Path(file_path).stem),
            }

        except Exception as e:
            logger.error(f"Error processing EPUB {file_path}: {e}")
            raise

    def process_arxiv(self, arxiv_id: str) -> Dict[str, Any]:
        """
        Download and process paper from arXiv.

        Args:
            arxiv_id: arXiv paper ID (e.g., "2301.12345")

        Returns:
            Dictionary containing text and metadata
        """
        try:
            # Search for the paper
            search = arxiv.Search(id_list=[arxiv_id])
            paper = next(search.results())

            # Download PDF
            filename = sanitize_filename(f"{arxiv_id}.pdf")
            file_path = self.document_dir / filename
            paper.download_pdf(str(file_path))

            # Process the downloaded PDF
            pdf_result = self.process_pdf(str(file_path))

            # Add arXiv metadata
            pdf_result["metadata"].update({
                "arxiv_id": arxiv_id,
                "title": paper.title,
                "authors": [author.name for author in paper.authors],
                "published": paper.published.isoformat(),
                "summary": paper.summary,
                "categories": paper.categories,
            })
            pdf_result["title"] = paper.title

            return {
                **pdf_result,
                "file_path": str(file_path),
            }

        except Exception as e:
            logger.error(f"Error processing arXiv paper {arxiv_id}: {e}")
            raise

    def save_document(
        self,
        file_path: str,
        document_type: Optional[DocumentType] = None,
        arxiv_id: Optional[str] = None
    ) -> Document:
        """
        处理并保存文档的核心方法

        Process and save a document.

        完整工作流程 (Complete Workflow):
        1. 根据文档类型调用相应的处理方法提取文本和元数据
        2. 生成唯一的文档ID(基于文件哈希值的前16位)
        3. 创建标准化的Document对象
        4. 保存两个文件:
           - {doc_id}.json: 文档元数据
           - {doc_id}_text.txt: 提取的纯文本

        文档ID生成策略 (Document ID Strategy):
        - 使用文件内容的哈希值确保唯一性
        - 相同内容的文件会有相同的ID(自动去重)
        - 取哈希值前16位作为ID(足够短且碰撞概率极低)

        Args:
            file_path: 文档文件路径 (Path to document file)
            document_type: 文档类型(None时自动检测) (Type of document)
            arxiv_id: arXiv论文ID(可选) (arXiv ID if processing arXiv paper)

        Returns:
            Document对象，包含所有元数据 (Document object with metadata)
        """
        # 根据类型处理文档
        if arxiv_id:
            # 处理arXiv论文(下载并提取)
            result = self.process_arxiv(arxiv_id)
            file_path = result["file_path"]
            document_type = DocumentType.ARXIV
        else:
            # 自动检测文档类型
            if document_type is None:
                document_type = get_document_type(file_path)
                if document_type is None:
                    raise ValueError(f"Unsupported file type: {file_path}")

            # 根据文档类型调用相应的处理方法
            if document_type == DocumentType.PDF:
                result = self.process_pdf(file_path)
            elif document_type == DocumentType.DOCX:
                result = self.process_docx(file_path)
            elif document_type == DocumentType.PPTX:
                result = self.process_pptx(file_path)
            elif document_type == DocumentType.EPUB:
                result = self.process_epub(file_path)
            else:  # TXT, MD
                result = self.process_text_file(file_path)

        # 生成文档唯一ID(使用文件哈希的前16位)
        file_hash = get_file_hash(file_path)
        doc_id = file_hash[:16]

        # 创建文档对象
        document = Document(
            id=doc_id,
            title=result["title"],
            file_path=file_path,
            document_type=document_type,
            file_size=get_file_size(file_path),
            num_pages=result.get("num_pages"),
        )

        # 保存元数据为JSON文件
        # 存储位置: data/metadata/{doc_id}.json
        metadata_file = self.metadata_dir / f"{doc_id}.json"
        metadata = {
            **document.model_dump(mode='json'),
            "extracted_metadata": result.get("metadata", {}),
        }

        with open(metadata_file, 'w', encoding='utf-8') as f:
            json.dump(metadata, f, indent=2, ensure_ascii=False, default=str)

        # 保存提取的文本为纯文本文件
        # 存储位置: data/metadata/{doc_id}_text.txt
        text_file = self.metadata_dir / f"{doc_id}_text.txt"
        with open(text_file, 'w', encoding='utf-8') as f:
            f.write(result["text"])

        logger.info(f"Saved document {doc_id}: {document.title}")
        return document

    def get_document_text(self, doc_id: str) -> Optional[str]:
        """
        Retrieve extracted text for a document.

        Args:
            doc_id: Document ID

        Returns:
            Extracted text or None if not found
        """
        text_file = self.metadata_dir / f"{doc_id}_text.txt"
        if text_file.exists():
            with open(text_file, 'r', encoding='utf-8') as f:
                return f.read()
        return None

    def get_document_metadata(self, doc_id: str) -> Optional[Document]:
        """
        Retrieve metadata for a document.

        Args:
            doc_id: Document ID

        Returns:
            Document object or None if not found
        """
        metadata_file = self.metadata_dir / f"{doc_id}.json"
        if metadata_file.exists():
            with open(metadata_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return Document(**data)
        return None

    def list_documents(self) -> List[Document]:
        """
        List all processed documents.

        Returns:
            List of Document objects
        """
        documents = []
        for metadata_file in self.metadata_dir.glob("*.json"):
            if not metadata_file.stem.endswith("_text"):
                with open(metadata_file, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    documents.append(Document(**data))
        return documents

    def delete_document(self, doc_id: str) -> bool:
        """
        Delete a document and its metadata.

        Args:
            doc_id: Document ID

        Returns:
            True if deleted successfully
        """
        try:
            metadata_file = self.metadata_dir / f"{doc_id}.json"
            text_file = self.metadata_dir / f"{doc_id}_text.txt"

            if metadata_file.exists():
                metadata_file.unlink()
            if text_file.exists():
                text_file.unlink()

            logger.info(f"Deleted document {doc_id}")
            return True
        except Exception as e:
            logger.error(f"Error deleting document {doc_id}: {e}")
            return False
